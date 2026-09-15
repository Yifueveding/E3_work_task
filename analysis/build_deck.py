"""Build the GCV Solar + Storage Valuation slide deck (11 slides)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- palette (matches chart styling) ----------
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK_SECONDARY = RGBColor(0x52, 0x51, 0x4E)
INK_MUTED = RGBColor(0x89, 0x87, 0x81)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
PAGE_BG = RGBColor(0xFF, 0xFF, 0xFF)
GRIDLINE = RGBColor(0xE1, 0xE0, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = 'Calibri'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW, SH = prs.slide_width, prs.slide_height


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=PAGE_BG):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = color
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             space_after=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return tb


def add_bullets(slide, x, y, w, h, items, size=15, color=INK_SECONDARY, font=FONT,
                 space_after=10, line_spacing=1.08, bullet_color=BLUE):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            txt, level = item
        else:
            txt, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        p.level = level
        marker = '—  ' if level == 0 else '·  '
        r = p.add_run()
        r.text = marker + txt
        r.font.size = Pt(size - (2 if level else 0))
        r.font.name = font
        r.font.color.rgb = color if level == 0 else INK_MUTED
    return tb


def slide_header(slide, kicker, title, n, total=11):
    add_rect(slide, 0, 0, SW, Inches(0.09), BLUE)
    add_text(slide, Inches(0.55), Inches(0.30), Inches(10.5), Inches(0.35),
              kicker.upper(), size=13, color=BLUE, bold=True)
    add_text(slide, Inches(0.55), Inches(0.62), Inches(11.8), Inches(0.7),
              title, size=27, color=INK, bold=True)
    add_rect(slide, Inches(0.55), Inches(1.32), Inches(1.0), Pt(3), BLUE)
    # footer
    add_text(slide, Inches(0.55), SH - Inches(0.42), Inches(8), Inches(0.3),
              'GCV Solar + Storage Valuation  |  E3 Applicant Exercise', size=9.5, color=INK_MUTED)
    add_text(slide, SW - Inches(1.2), SH - Inches(0.42), Inches(0.7), Inches(0.3),
              f'{n} / {total}', size=9.5, color=INK_MUTED, align=PP_ALIGN.RIGHT)


def add_picture_fit(slide, path, x, y, max_w, max_h):
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    px = x + int((max_w - w) / 2)
    py = y + int((max_h - h) / 2)
    slide.shapes.add_picture(path, px, py, width=w, height=h)


def stat_tile(slide, x, y, w, h, value, label, value_color=BLUE):
    add_rect(slide, x, y, w, h, SURFACE)
    box = slide.shapes[-1]
    box.line.color.rgb = GRIDLINE
    box.line.width = Pt(0.75)
    add_text(slide, x + Inches(0.15), y + Inches(0.15), w - Inches(0.3), Inches(0.55),
              value, size=26, color=value_color, bold=True)
    add_text(slide, x + Inches(0.15), y + h - Inches(0.55), w - Inches(0.3), Inches(0.45),
              label, size=11.5, color=INK_SECONDARY)


IMG = 'analysis/'  # image directory (run from repo root)

# =====================================================================
# SLIDE 1 — Title
# =====================================================================
s = add_slide()
set_bg(s, RGBColor(0x0B, 0x0B, 0x0B))
add_rect(s, 0, Inches(6.55), SW, Inches(0.95), BLUE)
add_text(s, Inches(0.9), Inches(2.35), Inches(11.5), Inches(1.6),
          'GCV Solar + Storage Valuation', size=42, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.6),
          'NYISO N.Y.C. Zone  |  2023 Backcast Revenue Analysis & Investment Recommendation',
          size=18, color=RGBColor(0xC3, 0xC2, 0xB7))
add_rect(s, Inches(0.9), Inches(4.25), Inches(0.9), Pt(3), BLUE)
add_text(s, Inches(0.9), Inches(6.72), Inches(8), Inches(0.5),
          'Prepared for Green Charge Ventures  |  E3 Applicant Exercise', size=13, color=WHITE, bold=True)
add_text(s, SW - Inches(2.4), Inches(6.72), Inches(1.9), Inches(0.5),
          '2026', size=13, color=WHITE, bold=True, align=PP_ALIGN.RIGHT)

# =====================================================================
# SLIDE 2 — Executive Summary
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Executive Summary', 'Combined solar + storage is the recommended installation', 2)

tiles = [
    ('$33.94/MWh', '2023 avg. NYISO N.Y.C.\nday-ahead price'),
    ('$18.8M', 'Standalone solar\n2023 backcast revenue'),
    ('$2.0M', 'Standalone storage\n2023 backcast revenue'),
    ('$20.8M', 'Combined system\n2023 backcast revenue'),
]
tw, th, gap = Inches(2.75), Inches(1.35), Inches(0.25)
x0 = Inches(0.55)
for i, (val, lab) in enumerate(tiles):
    stat_tile(s, x0 + i * (tw + gap), Inches(1.6), tw, th, val, lab)

add_bullets(s, Inches(0.55), Inches(3.35), Inches(6.0), Inches(3.4), [
    'Prices peak in winter (Feb, cold-snap heating demand) and again in summer (Jul, A/C peak); shoulder months (spring/fall) trough well below average.',
    'Solar earns slightly above the flat average because its output aligns with summer peak-price hours; it earns below average in winter.',
    'Storage value tracks price volatility, not price level — Jul and Feb (largest daily price swings) are its best months.',
    "Pairing solar + storage adds only a marginal energy-arbitrage benefit today (+0.05%) since inverter clipping is minimal, but the combined package's cheaper per-unit equipment pricing saves ~$35M in capex vs. building both standalone.",
], size=14.5)

add_rect(s, Inches(6.85), Inches(3.35), Inches(5.95), Inches(3.55), SURFACE)
box = s.shapes[-1]; box.line.color.rgb = BLUE; box.line.width = Pt(1.25)
add_text(s, Inches(7.1), Inches(3.55), Inches(5.4), Inches(0.4), 'RECOMMENDATION', size=12, color=BLUE, bold=True)
add_text(s, Inches(7.1), Inches(3.9), Inches(5.4), Inches(0.6), 'Combined Solar + Storage', size=21, color=INK, bold=True)
add_bullets(s, Inches(7.1), Inches(4.55), Inches(5.4), Inches(2.2), [
    '~21-year simple payback — matches standalone solar, so storage is effectively "added for free" from a payback standpoint.',
    '~$35M cheaper in capex than building solar and storage as two separate standalone projects.',
    "Better hedges the coming decade: storage's value should rise as solar buildout erodes solar's own capture price.",
], size=13.5)

# =====================================================================
# SLIDE 3 — Q1a Historical Prices
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q1a  ·  Historical Prices', '2023 Monthly Average Day-Ahead Energy Price', 3)
add_picture_fit(s, IMG + 'monthly_da_price_2023.png', Inches(0.4), Inches(1.55), Inches(7.7), Inches(5.3))
add_bullets(s, Inches(8.35), Inches(1.7), Inches(4.5), Inches(4.9), [
    'Annual average: $33.94/MWh.',
    'Winter peak: Feb ($51.84) and Jan ($42.77) — cold-snap driven heating & gas demand.',
    'Summer peak: Jul ($41.72) — A/C-driven peak demand.',
    'Shoulder-month trough: Apr–Jun & Aug–Oct sit near $25–30/MWh, when heating and cooling demand are both low.',
    'This "winter + summer peak, shoulder trough" shape drives which hours are most valuable for solar vs. storage (next slide).',
], size=14.5)

# =====================================================================
# SLIDE 4 — Q1b/Q1c Solar vs Storage seasonal value
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q1b & Q1c  ·  Seasonal Value', 'Solar and Storage Peak in Different — and Overlapping — Months', 4)
half_w = Inches(6.05)
add_picture_fit(s, IMG + 'solar_value_by_month.png', Inches(0.4), Inches(1.5), half_w, Inches(3.55))
add_picture_fit(s, IMG + 'storage_value_by_month.png', Inches(6.85), Inches(1.5), half_w, Inches(3.55))

add_text(s, Inches(0.4), Inches(5.15), half_w, Inches(0.35), 'Solar: most valuable May–Sep', size=13.5, color=BLUE, bold=True)
add_bullets(s, Inches(0.4), Inches(5.55), half_w, Inches(1.7), [
    'Summer capture price runs 8–17% above the flat average (Jul: $48.89 vs. $41.72) — solar hours line up with A/C-driven peak prices.',
    'Winter capture price runs below flat average (95–99%) — winter price spikes hit before sunrise / after sunset, which solar misses.',
], size=12.5)

add_text(s, Inches(6.85), Inches(5.15), half_w, Inches(0.35), 'Storage: most valuable in extreme-weather months', size=13.5, color=ORANGE, bold=True)
add_bullets(s, Inches(6.85), Inches(5.55), half_w, Inches(1.7), [
    'Jul ($50/MWh avg. daily spread) and Feb ($43/MWh) are far above the $29 annual average — driven by sharp heat-wave / cold-snap price spikes.',
    'Shoulder months (Mar–Jun, Aug–Dec) are weakest ($21–30) — mild weather means flatter daily price shapes and less arbitrage opportunity.',
], size=12.5)

# =====================================================================
# SLIDE 5 — Q2 Methodology & Assumptions
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q2  ·  Methodology', 'Revenue Modeling Approach & Key Assumptions', 5)

chips = [
    ('300 MW-DC / 250 MW-AC', 'Solar Size / Inverter Limit'),
    ('50 MW / 200 MWh (4-hr)', 'Storage Power / Energy'),
    ('100%', 'Round-Trip Efficiency'),
    ('Day-Ahead (DA)', 'Price Basis — see note below'),
]
chip_w = Inches((12.2 - 3 * 0.15) / 4)
chip_h = Inches(0.72)
chip_y = Inches(1.42)
for i, (val, lab) in enumerate(chips):
    cx = Inches(0.55) + i * (chip_w + Inches(0.15))
    add_rect(s, cx, chip_y, chip_w, chip_h, SURFACE)
    box = s.shapes[-1]; box.line.color.rgb = GRIDLINE; box.line.width = Pt(0.75)
    add_text(s, cx + Inches(0.12), chip_y + Inches(0.09), chip_w - Inches(0.24), Inches(0.32),
              val, size=13.5, color=BLUE, bold=True)
    add_text(s, cx + Inches(0.12), chip_y + Inches(0.42), chip_w - Inches(0.24), Inches(0.28),
              lab, size=9.5, color=INK_SECONDARY)

add_text(s, Inches(0.55), Inches(2.24), Inches(12.2), Inches(0.24),
          'Notation: p_t = hourly price ($/MWh); s_t = normalized solar shape; G_t = delivered solar MW.',
          size=10, color=INK_MUTED, italic=True)

rows = [
    ('Standalone Solar', 'Delivered output = min(Solar Shape × 300 MW-DC, 250 MW-AC inverter limit). Revenue = Σ(delivered MW × DA price).', 0.5),
    ('Standalone Storage', 'Per-day linear program (50 MW power limit, 200 MWh energy limit, SOC starts at 0) chooses the revenue-maximizing hourly charge/discharge schedule — letting the optimizer find multiple cycles/day when DA prices oscillate enough, rather than assuming one fixed 4-hr cycle (worth +14.75% vs. that heuristic). Perfect day-ahead price foresight; RTE = 100% (per given input); no degradation/cycling cost; energy arbitrage only in this base case (see Reserve Scenario row below).', 1.0),
    ('Combined Solar + Storage', 'Solar identical to standalone. Storage uses the same per-day LP, except charging can draw up to 50 MW/hr for $0 from that hour\'s otherwise-clipped solar before paying the grid price for any remaining need.', 0.65),
    ('Reserve Scenario (Q2d)', 'Extension beyond the base exercise: reserve capacity valued at $1.25/MW-hr (midpoint of $1–1.5/MW). Headroom = up-reserve capacity; footroom = down-reserve capacity, classified hour-by-hour from the actual LP dispatch. Solar cannot provide reserve (variable, non-dispatchable) → $0. Storage → +$0.78M — less than a single-cycle heuristic would offer, since multi-cycling leaves fewer idle hours available as reserve. Details on Slide 8.', 0.95),
]

ty = Inches(2.56)
col1_w, col2_w = Inches(2.9), Inches(9.35)
add_rect(s, Inches(0.55), ty, col1_w + col2_w, Pt(2), GRIDLINE)
ry = ty + Inches(0.12)
for label, desc, h in rows:
    row_h = Inches(h)
    add_text(s, Inches(0.55), ry, col1_w, Inches(0.4), label, size=13.5, color=BLUE, bold=True)
    add_text(s, Inches(3.55), ry, col2_w, row_h, desc, size=11, color=INK_SECONDARY, line_spacing=1.08)
    ry = ry + row_h + Inches(0.12)
    add_rect(s, Inches(0.55), ry - Inches(0.08), col1_w + col2_w, Pt(1), GRIDLINE)

add_text(s, Inches(0.55), ry + Inches(0.06), Inches(12.2), Inches(0.6),
          'All simplifications are intentional given the exercise\'s time-box (Q2b/2c each capped at ~1 hour); more rigorous operating assumptions are proposed in Q3b. Note: we use DA prices specifically because the LP dispatch assumes perfect price foresight — realistic under DA (published in advance) but not under real-time (not knowable in advance). A parallel RT-basis run earns modestly more (solar $18.91M, storage $2.73M) precisely because it credits foreknowledge an RT-only strategy wouldn\'t really have.',
          size=11, color=INK_MUTED, italic=True)

# =====================================================================
# SLIDE 6 — Q2 Results
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q2  ·  Results', '2023 Backcast Revenue by Configuration', 6)
add_picture_fit(s, IMG + 'q2_revenue_by_config.png', Inches(0.4), Inches(1.5), Inches(7.9), Inches(5.4))
add_bullets(s, Inches(8.5), Inches(1.7), Inches(4.35), Inches(4.9), [
    'Standalone Solar: $18.80M (551,454 MWh delivered; avg. capture price $34.10/MWh).',
    'Standalone Storage: $1.98M ($39.56/kW-yr, $9.89/kWh-yr) — multi-cycle LP-optimized energy arbitrage only.',
    'Combined System: $20.79M — essentially the sum of the two standalone revenues.',
    'Solar dominates the revenue stack for all three configurations; storage revenue from energy arbitrage alone is comparatively small.',
], size=14)

# =====================================================================
# SLIDE 7 — Q2c Interaction Detail
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q2c  ·  Solar + Storage Interaction', 'The Pairing Benefit Is Real but Small Today', 7)

add_bullets(s, Inches(0.55), Inches(1.7), Inches(6.1), Inches(4.5), [
    'In the combined system, storage can charge from solar output that would otherwise be clipped by the 250 MW inverter limit — at zero marginal cost.',
    'Gross DC solar generation: 551,826 MWh; clipped/curtailed: only 372 MWh (0.07% of gross) — the 300/250 MW (1.2x) DC:AC ratio isn\'t aggressive enough to cause meaningful clipping.',
    'Combined revenue ($20,791,797) exceeds the simple sum of standalone solar + standalone storage ($20,781,773) by just $10,023 (+0.05%).',
], size=15, space_after=16)

add_rect(s, Inches(7.1), Inches(1.7), Inches(5.7), Inches(3.0), SURFACE)
box = s.shapes[-1]; box.line.color.rgb = GRIDLINE; box.line.width = Pt(1)
add_text(s, Inches(7.35), Inches(1.9), Inches(5.2), Inches(0.4), 'TAKEAWAY', size=12, color=BLUE, bold=True)
add_text(s, Inches(7.35), Inches(2.3), Inches(5.2), Inches(2.2),
          'On this design, pairing solar and storage does not meaningfully change energy revenue vs. running them separately.\n\nThe real case for the combined system rests on lower blended equipment cost, not an energy-arbitrage synergy — see Q3c.',
          size=14, color=INK_SECONDARY, line_spacing=1.25)

add_text(s, Inches(0.55), Inches(6.3), Inches(12.2), Inches(0.6),
          'Looking ahead: as NY\'s solar buildout accelerates over the next decade, curtailment risk — and therefore this interaction benefit — should grow (see Q3a).',
          size=13, color=INK_MUTED, italic=True)

# =====================================================================
# SLIDE 8 — Q2d Reserve (Headroom/Footroom) Scenario [extension, beyond base exercise]
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q2d  ·  Reserve Scenario (Extension)', 'Layering NYISO Reserve Revenue on Top of Energy Arbitrage', 8)

half_w = Inches(6.05)
add_picture_fit(s, IMG + 'q2_energy_only_by_config.png', Inches(0.4), Inches(1.5), half_w, Inches(3.55))
add_picture_fit(s, IMG + 'q2_energy_plus_reserve_by_config.png', Inches(6.85), Inches(1.5), half_w, Inches(3.55))

add_bullets(s, Inches(0.55), Inches(5.15), Inches(12.2), Inches(1.5), [
    'Beyond the base exercise: estimates additional revenue from offering NYISO reserve capacity at $1.25/MW-hr (midpoint of the $1–1.5/MW range you flagged). Headroom = capacity to increase output on call; footroom = capacity to decrease output on call.',
    'Solar cannot provide reserve — it is a variable, non-dispatchable resource, so it earns $0. Storage is classified hour-by-hour from its actual Q2b/c LP dispatch (footroom while charging, headroom while discharging, both simultaneously while idle) → +$0.78M (+39.5%) — less than a single-cycle heuristic would offer, since multi-cycling leaves fewer idle hours. Combined reserve comes entirely from the storage component → +$0.78M (+3.8%).',
], size=12.5, space_after=6)

add_text(s, Inches(0.55), Inches(6.72), Inches(12.2), Inches(0.3),
          'Illustrative only — real reserve products carry stricter rules than modeled here; see Q3b for a rigorous treatment.',
          size=12, color=INK_MUTED, italic=True)

# =====================================================================
# SLIDE 9 — Q3a Revenue Outlook
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q3a  ·  Forward Outlook', 'How Will Each Revenue Stream Evolve Over the Next Decade?', 9)

cols = [
    ('Standalone Solar', 'Capture price likely erodes', [
        "CLCPA mandates 70% renewables by 2030, 100% zero-emission grid by 2040 — a large wave of new solar enters NYISO.",
        'More solar depresses midday prices ("duck curve" / cannibalization), even as the overall price level may rise with electrification-driven demand.',
    ], BLUE),
    ('Standalone Storage', 'Value likely increases', [
        'Same trend that hurts solar helps storage: steeper evening ramps widen the daily price spread storage arbitrages.',
        "NYC's Peaker Rule retirements tighten local capacity. Risk: storage cannibalization if buildout outpaces the opportunity.",
    ], ORANGE),
    ('Combined System', 'Interaction benefit should grow', [
        "Today's interaction (+0.05%) reflects minimal clipping.",
        'As solar penetration rises, a co-located battery\'s ability to absorb otherwise-curtailed solar becomes more valuable — strengthening the strategic case for pairing.',
    ], INK),
]
cw = Inches(3.95); gap = Inches(0.2); x0 = Inches(0.55)
for i, (title, sub, bullets, color) in enumerate(cols):
    x = x0 + i * (cw + gap)
    add_rect(s, x, Inches(1.6), cw, Inches(0.08), color)
    add_text(s, x, Inches(1.85), cw, Inches(0.5), title, size=17, color=INK, bold=True)
    add_text(s, x, Inches(2.35), cw, Inches(0.5), sub, size=13.5, color=color, bold=True, italic=True)
    add_bullets(s, x, Inches(2.95), cw, Inches(3.8), bullets, size=13, space_after=14, line_spacing=1.15)

# =====================================================================
# SLIDE 10 — Q3b Follow-on Study
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q3b  ·  Proposed Follow-On Study', 'A More Complete Revenue Forecast — 10 Additional Hours', 10)

items = [
    ('1', 'Forward price shapes', '~3 hrs', 'Replace the 2023 backcast with a forward-looking hourly price shape reflecting expected capacity additions/retirements and load growth (electrification, data centers), grounded in published NYISO / NYSERDA CLCPA outlooks.'),
    ('2', 'Realistic operating assumptions', '~2 hrs', 'Round-trip efficiency ~85–90% (not 100%), solar degradation (~0.5%/yr), battery degradation/augmentation cost. Q2b already uses a per-day LP that finds multi-cycle dispatch worth +14.75% under a zero-cost assumption — adding a realistic per-cycle degradation cost would temper that toward an economically optimal cycling frequency.'),
    ('3', 'Capacity & ancillary revenue', '~2 hrs', "Add NYISO capacity market (ICAP) and frequency regulation revenue — often the majority of a merchant battery's revenue stack. Q2d sketches an illustrative reserve add-on; a rigorous version (real product rules, co-optimization, capacity market) is still the single biggest gap."),
    ('4', 'Scenario / sensitivity analysis', '~2 hrs', 'High/low cases on load growth, gas prices, and renewable build-out pace, given how uncertain the 10-year price path is.'),
    ('5', 'Full pro forma', '~1 hr', 'Layer in provided capex, O&M, financing, and ITC/IRA tax credit treatment to compute NPV/IRR rather than simple payback.'),
]
ty = Inches(1.65)
row_h = Inches(0.98)
for i, (num, title, hrs, desc) in enumerate(items):
    ry = ty + i * row_h
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), ry, Inches(0.5), Inches(0.5))
    circ.fill.solid(); circ.fill.fore_color.rgb = BLUE; circ.line.fill.background(); circ.shadow.inherit = False
    tf = circ.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    add_text(s, Inches(1.25), ry - Inches(0.02), Inches(3.1), Inches(0.4), title, size=15, color=INK, bold=True)
    add_rect(s, Inches(4.45), ry + Inches(0.02), Inches(0.85), Inches(0.32), SURFACE)
    box = s.shapes[-1]; box.line.color.rgb = BLUE; box.line.width = Pt(0.75)
    add_text(s, Inches(4.45), ry + Inches(0.06), Inches(0.85), Inches(0.28), hrs, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(5.5), ry - Inches(0.02), Inches(7.3), Inches(0.85), desc, size=12, color=INK_SECONDARY, line_spacing=1.1)
    if i < len(items) - 1:
        add_rect(s, Inches(0.55), ry + row_h - Inches(0.12), Inches(12.25), Pt(0.75), GRIDLINE)

# =====================================================================
# SLIDE 11 — Q3c Recommendation
# =====================================================================
s = add_slide(); set_bg(s)
slide_header(s, 'Q3c  ·  Recommendation', 'Combined Solar + Storage Is the Recommended Installation', 11)

add_picture_fit(s, IMG + 'q3_payback_comparison.png', Inches(0.4), Inches(1.5), Inches(7.7), Inches(3.5))

add_rect(s, Inches(0.4), Inches(5.15), Inches(7.7), Inches(0.05), GRIDLINE)
add_bullets(s, Inches(0.4), Inches(5.3), Inches(7.7), Inches(1.9), [
    'Simple payback (capex ÷ 2023 backcast revenue, no financing/tax/O&M): Standalone Solar ~22 yrs, Standalone Storage ~38 yrs, Combined ~21 yrs.',
], size=12.5, color=INK_MUTED)

add_rect(s, Inches(8.35), Inches(1.5), Inches(4.5), Inches(5.35), SURFACE)
box = s.shapes[-1]; box.line.color.rgb = BLUE; box.line.width = Pt(1.25)
add_text(s, Inches(8.6), Inches(1.7), Inches(4.0), Inches(0.4), 'WHY COMBINED', size=12, color=BLUE, bold=True)
add_bullets(s, Inches(8.6), Inches(2.15), Inches(4.05), Inches(3.2), [
    "Closely matches (and slightly beats) standalone solar's payback — storage is effectively added at no economic cost.",
    'Combined package pricing ($1.25/W solar, $350/kWh storage) saves ~$35M in capex vs. building both standalone.',
    'Better hedges the Q3a outlook: storage value should rise as solar capture price erodes.',
], size=12.5, space_after=12)
add_text(s, Inches(8.6), Inches(5.35), Inches(4.05), Inches(1.6),
          'Caveat: standalone storage still looks weak here because only day-ahead energy arbitrage was modeled. Real storage revenue also includes capacity & ancillary services (Q3b) — revisit before committing capital.',
          size=11.5, color=INK_MUTED, italic=True, line_spacing=1.2)

prs.save('GCV_Solar_Storage_Valuation.pptx')
print('Saved GCV_Solar_Storage_Valuation.pptx with', len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
